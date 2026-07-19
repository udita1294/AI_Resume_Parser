import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")
if not my_api_key:
    raise ValueError("GROQ_API_KEY not found.")

client = Groq(api_key=my_api_key)
model = "llama-3.3-70b-versatile"

job_description = """
    Fullstack Developer :

We are seeking an experienced and highly motivated
Fullstack Developer to join our product development team. The ideal candidate
will have a strong background in both front-end and back-end development, with
a focus on building and maintaining scalable, high-performance applications.
This role requires a proven track record in product development, from concept
to deployment.

Roles and Responsibilities :
1) Responsibilities Design, develop, and maintain robust and
scalable full-stack applications
2) Develop user-facing features using ReactJS.
Build efficient and relible server-side components using either Java
Springboot or NodeJS.
3) Manage and optimize database interactions, utilizing
both Relational and Non-Relational databases.
4) Collaborate with Product Managers, designers, and other engineers to define,
design, and ship new features.
5) Participate in code reviews to ensure code quality and adherence to
best practices.
6) Troubleshoot, debug, and upgrade existing systems. Contribute
to the entire application lifecycle, from conception to deployment. Required

Qualifications:
0-1 years of professional experience in Fullstack Development.
Strong proficiency in front-end development using ReactJS and its ecosystem
(Redux, Hooks, etc.). Demonstrated experience in back-end development with
either Java Springboot or NodeJS. Solid understanding and practical
experience with database technologies, including both Relational Databases
(e.g., PostgreSQL, MySQL) and Non-Relational Databases (e.g., MongoDB,
DynamoDB).
Mandatory experience in the full lifecycle of Product Development, from
initial requirement gathering to deployment and maintenance. Experience with
RESTful APIs and asynchronous request handling. Familiarity with version
control systems (e.g., Git). Preferred Qualifications Experience with cloud
platforms (AWS, Azure, GCP). Familiarity with containerization and
orchestration tools (Docker, Kubernetes). Knowledge of modern authorization
mechanisms, such as JSON Web Tokens (JWT). Experience working in an Agile/Scrum development environment.
"""

class JobDescription(BaseModel):
    role:str
    required_skills:list[str]
    preferred_skills:list[str]
    minimum_experience:float|None
    education_requirements:list[str]
    responsibilities:list[str]

jobD_schema = JobDescription.model_json_schema()

system_prompt = f"""
    You are an expert HR Assistant.
    Your job is to analyse the job descriptions and extract structured informations from them.

    Return ONLY valid JSON matching this schema:
    {jobD_schema}

    IMPORTANT:
    - Do not return the schema itself.
    - Do NOT return fields like "properties","title" or "type".
    - Fill the schema with the actual information extracted by the job description.

    If minimum experience is not mentioned in the job description, return null for that field.
    If information for a list is missing return an empty list for that field.
    Do NOT invent any information that is not present in the job description.
"""

user_prompt = f"""
    Analyze the following job description.
    {job_description}
"""

message_system = {
    "role": "system",
    "content": system_prompt
}

message_user = {
    "role": "user",
    "content": user_prompt
}

response_format = {
    "type": "json_object"
}

messages = [message_system, message_user]

response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
answer = response.choices[0].message.content
raw_json = answer
# print(raw_json)

import json
job_data = json.loads(raw_json)
job = JobDescription(**job_data)

print("Role:", job.role)
print("Minimum Experience:", job.minimum_experience)
print("Required Skills:", job.required_skills)


# parse real
class MatchResult(BaseModel):
    score : float
    details : dict

class Experience(BaseModel):
    company : str | None = None
    role : str | None = None
    duration : str | None = None
    description : str | None = None
    skill_used : list[str] = []

class Resume(BaseModel):
    name : str | None = None
    email : str | None = None
    phone : str | None = None
    total_experience_years : float | None = None
    skills : list[str] = []
    experiences : list[Experience] = []
    education : list[str] = []
    projects : list[str] = []
    certifications : list[str] = []

resume_schema = Resume.model_json_schema()

def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.
    Commpare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}
    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:
    {match_schema}

    Give me :
    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether the experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response consise to read.
"""
    
    message ={
    "role": "user",
    "content": prompt
    }
    messages = [message]
    response_format = {
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)

def parse_resume(resume_text):
    system_prompt = f"""
      You are an expert resume parser.
      Extract information from the resume based on its meaning, not only based on exact section headings.

      Different resumes may use different formats and section names, so you need to understand the context and extract relevant information accordingly.

      For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """

    user_prompt = f"""
    Parse the following resume:
    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


from pypdf import PdfReader
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text

